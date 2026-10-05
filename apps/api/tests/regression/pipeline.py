"""Runs the regression farms Excel through the same steps the app does.

1. Read the first sheet like the frontend does (`apps/web/src/utils/excel.ts`,
   `loadExcelFileFarmsData`): headers on row 2, data from row 4, headers matched by
   keyword in Spanish or English. The upload has no country column: like the
   frontend, every row gets the analysis country.
2. `POST /farms/parse`, then `POST /deforestation_analysis/analize` against every
   layer in the index, and `POST /polygons_validation/validate`.
3. Summarize what matters for regressions: parsed areas, one deforestation ratio per
   farm and layer, validation statuses and inconsistencies.
"""

import math
from contextlib import nullcontext
from datetime import date, datetime
from pathlib import Path
from typing import Any
from unittest.mock import patch

import openpyxl
from fastapi.testclient import TestClient

from app.main import app
from app.modules.maps.helpers import get_all_maps

REGRESSION_DIR = Path(__file__).parent
EXCEL_PATH = REGRESSION_DIR / "regression_farms.xlsx"
EXPECTED_PATH = REGRESSION_DIR / "expected_results.json"
FIXTURE_MAPS_ROOT = REGRESSION_DIR / "fixtures" / "maps"

# Pinned so results don't depend on the local .env: the value deployed in Azure.
OVERLAP_THRESHOLD_PERCENTAGE = 1.0
LOCALE = "es"
# The frontend sets the country chosen on its landing page on every row. The farms
# are in several countries, but no result depends on it.
COUNTRY = "EC"

# attribute: (es-header, en-header), as in the frontend's headerKeywordsMappings
HEADER_KEYWORDS = {
    "id": ("id", "id"),
    "producerName": ("nombre productor", "producer name"),
    "productionDate": ("fecha producción", "production date"),
    "productionQuantity": ("cantidad producción", "production quantity"),
    "productionQuantityUnit": (
        "unidad cantidad producción",
        "production measurement unit",
    ),
    "region": ("región", "region"),
    "coordinatesFormat": ("formato coordenadas", "coordinates format"),
    "geometryType": ("tipo geometría", "geometry type"),
    "farmCoordinates": ("coordenadas finca", "land coordinates"),
    "area": ("superficie [hectáreas]", "area [hectares]"),
    "cropType": ("tipo de cultivo", "crop type"),
    "association": ("asociación", "cooperative"),
    "documentName1": ("nombre documento 1", "document name 1"),
    "documentUrl1": ("enlace documento 1", "document link 1"),
    "documentName2": ("nombre documento 2", "document name 2"),
    "documentUrl2": ("enlace documento 2", "document link 2"),
    "documentName3": ("nombre documento 3", "document name 3"),
    "documentUrl3": ("enlace documento 3", "document link 3"),
}


def read_farm_rows(path: Path = EXCEL_PATH) -> list[dict[str, Any]]:
    """The rows the frontend would send to `/farms/parse` for this Excel."""
    sheet = openpyxl.load_workbook(path, data_only=True).worksheets[0]

    key_by_column: dict[int, str] = {}
    for column in range(1, sheet.max_column + 1):
        header = sheet.cell(2, column).value
        if not header:
            continue
        parts = [part.strip().lower() for part in str(header).splitlines()]
        for key, keywords in HEADER_KEYWORDS.items():
            if any(keyword in parts for keyword in keywords):
                key_by_column[column] = key

    rows = []
    for row_number in range(4, sheet.max_row + 1):
        row: dict[str, Any] = {}
        for column, key in key_by_column.items():
            value = sheet.cell(row_number, column).value
            if isinstance(value, str):
                value = value.strip() or None
            elif isinstance(value, (datetime, date)):
                value = value.isoformat()
            row[key] = value
        if all(value is None for value in row.values()):
            continue
        if isinstance(row.get("id"), (int, float)):
            row["id"] = str(row["id"])
        row["country"] = COUNTRY
        row["documents"] = [
            {"name": row.get(f"documentName{i}") or "", "url": row[f"documentUrl{i}"]}
            for i in (1, 2, 3)
            if row.get(f"documentUrl{i}")
        ]
        for i in (1, 2, 3):
            row.pop(f"documentName{i}", None)
            row.pop(f"documentUrl{i}", None)
        rows.append(row)
    return rows


def run_pipeline(
    rows: list[dict[str, Any]],
    client=None,
    map_ids: list[int] | None = None,
    country: str | None = None,
) -> dict[str, Any]:
    """Parse, analyze and validate.

    By default it runs in-process against every layer in the current store, with
    the overlap threshold pinned. Pass an HTTP client (e.g. `httpx.Client` with a
    `base_url`) and the map ids to run it against a deployed API instead; the
    deployment's own threshold then applies. With the per-country layout, pass the
    `country` of the layers: ids are numbered within each country.
    """
    in_process = client is None
    if client is None:
        client = TestClient(app)

    response = client.post(f"/farms/parse?locale={LOCALE}", json=rows)
    response.raise_for_status()
    farms = response.json()
    payload = [
        {
            "id": farm["id"],
            "type": farm["polygon"]["type"],
            "details": farm["polygon"]["details"],
        }
        for farm in farms
    ]

    if map_ids is None:
        map_ids = sorted(
            entry["id"]
            for entry in get_all_maps()
            if country is None or entry["country"] in (None, country)
        )
    analysis_body: dict[str, Any] = {"maps": map_ids, "farms": payload}
    if country is not None:
        # Per-country layout: layer ids are numbered within each country.
        analysis_body["country"] = country
    response = client.post("/deforestation_analysis/analize", json=analysis_body)
    response.raise_for_status()
    analysis = response.json()

    pinned_threshold = (
        patch(
            "app.modules.polygons_validation.helpers.OVERLAP_THRESHOLD_PERCENTAGE",
            OVERLAP_THRESHOLD_PERCENTAGE,
        )
        if in_process
        else nullcontext()
    )
    with pinned_threshold:
        response = client.post("/polygons_validation/validate", json=payload)
    response.raise_for_status()
    validation = response.json()

    return {
        "farms": {
            farm["id"]: {
                "type": farm["polygon"]["type"],
                "area_m2": farm["polygon"]["area"],
            }
            for farm in farms
        },
        "deforestation": {
            farm["id"]: {
                str(layer["mapId"]): next(
                    result["value"]
                    for result in layer["farmResults"]
                    if result["farmId"] == farm["id"]
                )
                for layer in analysis
            }
            for farm in farms
        },
        "validation": {
            "status": {
                result["farmId"]: result["status"]
                for result in validation["farmResults"]
            },
            "inconsistencies": [
                {
                    "type": inconsistency["type"],
                    "farmIds": inconsistency["farmIds"],
                    **{
                        key: inconsistency["data"][key]
                        for key in ("percentage", "criticality", "area", "reason")
                        if key in inconsistency["data"]
                    },
                }
                for inconsistency in validation["inconsistencies"]
            ],
        },
    }


def differences(actual, expected, path="results") -> list[str]:
    """Human-readable list of every value that changed (floats compared to 1e-9)."""
    if isinstance(expected, dict) and isinstance(actual, dict):
        diffs = []
        for key in sorted(set(expected) | set(actual)):
            if key not in actual:
                diffs.append(f"{path}.{key}: missing (expected {expected[key]!r})")
            elif key not in expected:
                diffs.append(f"{path}.{key}: unexpected {actual[key]!r}")
            else:
                diffs += differences(actual[key], expected[key], f"{path}.{key}")
        return diffs
    if isinstance(expected, list) and isinstance(actual, list):
        if len(expected) != len(actual):
            return [f"{path}: {len(actual)} items, expected {len(expected)}"]
        return [
            diff
            for i, (a, e) in enumerate(zip(actual, expected))
            for diff in differences(a, e, f"{path}[{i}]")
        ]
    if isinstance(expected, float) and isinstance(actual, (int, float)):
        if math.isclose(actual, expected, rel_tol=1e-9, abs_tol=1e-12):
            return []
    elif actual == expected:
        return []
    return [f"{path}: {actual!r} (expected {expected!r})"]
