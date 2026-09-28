"""Runs the regression farms Excel through the same steps the app does.

1. Read the first sheet like the frontend does (`monbo-front/src/utils/excel.ts`,
   `loadExcelFileFarmsData`): headers on row 2, data from row 4, headers matched by
   keyword in Spanish or English.
2. `POST /farms/parse`, then `POST /deforestation_analysis/analize` against every
   layer in the index, and `POST /polygons_validation/validate`.
3. Summarize what matters for regressions: parsed areas, one deforestation ratio per
   farm and layer, validation statuses and inconsistencies.
"""

from datetime import date, datetime
from pathlib import Path
from typing import Any
from unittest.mock import patch

import openpyxl
from fastapi.testclient import TestClient

from app.main import app
from app.modules.maps.helpers import get_all_maps

REGRESSION_DIR = Path(__file__).parent
EXCEL_PATH = REGRESSION_DIR / "regression_farms.xlsx"
EXPECTED_PATH = REGRESSION_DIR / "expected_results.json"
FIXTURE_MAPS_ROOT = REGRESSION_DIR / "fixtures" / "maps"

# Pinned so results don't depend on the local .env: the value deployed in Azure.
OVERLAP_THRESHOLD_PERCENTAGE = 1.0
LOCALE = "es"

# attribute: (es-header, en-header), as in the frontend's headerKeywordsMappings
HEADER_KEYWORDS = {
    "id": ("id", "id"),
    "producerName": ("nombre productor", "producer name"),
    "productionDate": ("fecha producción", "production date"),
    "productionQuantity": ("cantidad producción", "production quantity"),
    "productionQuantityUnit": (
        "unidad cantidad producción",
        "production measurement unit",
    ),
    "country": ("país", "country"),
    "region": ("región", "region"),
    "coordinatesFormat": ("formato coordenadas", "coordinates format"),
    "geometryType": ("tipo geometría", "geometry type"),
    "farmCoordinates": ("coordenadas finca", "land coordinates"),
    "area": ("superficie [hectáreas]", "area [hectares]"),
    "cropType": ("tipo de cultivo", "crop type"),
    "association": ("asociación", "cooperative"),
    "documentName1": ("nombre documento 1", "document name 1"),
    "documentUrl1": ("enlace documento 1", "document link 1"),
    "documentName2": ("nombre documento 2", "document name 2"),
    "documentUrl2": ("enlace documento 2", "document link 2"),
    "documentName3": ("nombre documento 3", "document name 3"),
    "documentUrl3": ("enlace documento 3", "document link 3"),
}


def read_farm_rows(path: Path = EXCEL_PATH) -> list[dict[str, Any]]:
    """The rows the frontend would send to `/farms/parse` for this Excel."""
    sheet = openpyxl.load_workbook(path, data_only=True).worksheets[0]

    key_by_column: dict[int, str] = {}
    for column in range(1, sheet.max_column + 1):
        header = sheet.cell(2, column).value
        if not header:
            continue
        parts = [part.strip().lower() for part in str(header).splitlines()]
        for key, keywords in HEADER_KEYWORDS.items():
            if any(keyword in parts for keyword in keywords):
                key_by_column[column] = key

    rows = []
    for row_number in range(4, sheet.max_row + 1):
        row: dict[str, Any] = {}
        for column, key in key_by_column.items():
            value = sheet.cell(row_number, column).value
            if isinstance(value, str):
                value = value.strip() or None
            elif isinstance(value, (datetime, date)):
                value = value.isoformat()
            row[key] = value
        if all(value is None for value in row.values()):
            continue
        if isinstance(row.get("id"), (int, float)):
            row["id"] = str(row["id"])
        row["documents"] = [
            {"name": row.get(f"documentName{i}") or "", "url": row[f"documentUrl{i}"]}
            for i in (1, 2, 3)
            if row.get(f"documentUrl{i}")
        ]
        for i in (1, 2, 3):
            row.pop(f"documentName{i}", None)
            row.pop(f"documentUrl{i}", None)
        rows.append(row)
    return rows


def run_pipeline(rows: list[dict[str, Any]]) -> dict[str, Any]:
    """Parse, analyze against every layer in the current store, and validate."""
    client = TestClient(app)

    response = client.post(f"/farms/parse?locale={LOCALE}", json=rows)
    response.raise_for_status()
    farms = response.json()
    payload = [
        {
            "id": farm["id"],
            "type": farm["polygon"]["type"],
            "details": farm["polygon"]["details"],
        }
        for farm in farms
    ]

    map_ids = sorted(entry["id"] for entry in get_all_maps())
    response = client.post(
        "/deforestation_analysis/analize", json={"maps": map_ids, "farms": payload}
    )
    response.raise_for_status()
    analysis = response.json()

    with patch(
        "app.modules.polygons_validation.helpers.OVERLAP_THRESHOLD_PERCENTAGE",
        OVERLAP_THRESHOLD_PERCENTAGE,
    ):
        response = client.post("/polygons_validation/validate", json=payload)
    response.raise_for_status()
    validation = response.json()

    return {
        "farms": {
            farm["id"]: {
                "type": farm["polygon"]["type"],
                "area_m2": farm["polygon"]["area"],
            }
            for farm in farms
        },
        "deforestation": {
            farm["id"]: {
                str(layer["mapId"]): next(
                    result["value"]
                    for result in layer["farmResults"]
                    if result["farmId"] == farm["id"]
                )
                for layer in analysis
            }
            for farm in farms
        },
        "validation": {
            "status": {
                result["farmId"]: result["status"]
                for result in validation["farmResults"]
            },
            "inconsistencies": [
                {
                    "type": inconsistency["type"],
                    "farmIds": inconsistency["farmIds"],
                    **{
                        key: inconsistency["data"][key]
                        for key in ("percentage", "criticality", "area", "reason")
                        if key in inconsistency["data"]
                    },
                }
                for inconsistency in validation["inconsistencies"]
            ],
        },
    }
