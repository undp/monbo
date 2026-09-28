"""Regenerate the regression fixtures and expected results from the real layers.

    uv run python -m tests.regression.generate

Run it on purpose, when a change is *meant* to alter results (a new layer, a fixed
bug), and review the diff of `expected_results.json` in the pull request. It needs
the real rasters (`git lfs pull`).

It writes, next to this file:
- `fixtures/maps/`: a copy of the layers index plus one sparse GeoTIFF per layer.
  Each has the real raster's size, CRS and transform, but only the pixels around the
  regression farms are stored, so the fixtures stay small enough for Git (and CI,
  which doesn't fetch LFS) while giving exactly the same results as the real layers.
- `expected_results.json`: what `test_regression.py` compares against.
- the "Resultados esperados" sheet of `regression_farms.xlsx`, to compare by hand in
  the app.
"""

import json
import math
import shutil
from pathlib import Path

import openpyxl
import rasterio
from fastapi.testclient import TestClient
from rasterio.warp import transform_geom
from rasterio.windows import Window, from_bounds
from shapely.geometry import mapping, shape

from app.main import app
from app.models.farms import FarmPolygonDetailData
from app.modules.layers.store import LayerStore, get_layer_store, set_layer_store
from app.utils.farms import get_farm_coords_and_radius
from app.utils.polygons import generate_polygon
from tests.regression.pipeline import (
    EXCEL_PATH,
    EXPECTED_PATH,
    FIXTURE_MAPS_ROOT,
    LOCALE,
    OVERLAP_THRESHOLD_PERCENTAGE,
    read_farm_rows,
    run_pipeline,
)

PAD_PIXELS = 3
RESULTS_SHEET = "Resultados esperados"


def is_lfs_pointer(path: Path) -> bool:
    with open(path, "rb") as file:
        return file.read(40).startswith(b"version https://git-lfs")


def farm_polygons(rows):
    """The farm geometries exactly as the analysis builds them (points -> circles)."""
    response = TestClient(app).post(f"/farms/parse?locale={LOCALE}", json=rows)
    response.raise_for_status()
    polygons = []
    for farm in response.json():
        data = FarmPolygonDetailData(
            id=farm["id"],
            type=farm["polygon"]["type"],
            details=farm["polygon"]["details"],
        )
        coords, radius = get_farm_coords_and_radius(data)
        polygons.append(generate_polygon(coords, radius))
    return polygons


def padded_window(src, polygon) -> Window | None:
    left, bottom, right, top = shape(
        transform_geom("EPSG:4326", src.crs, mapping(polygon))
    ).bounds
    window = from_bounds(left, bottom, right, top, transform=src.transform)
    col = max(0, math.floor(window.col_off) - PAD_PIXELS)
    row = max(0, math.floor(window.row_off) - PAD_PIXELS)
    col_end = min(src.width, math.ceil(window.col_off + window.width) + PAD_PIXELS)
    row_end = min(src.height, math.ceil(window.row_off + window.height) + PAD_PIXELS)
    if col_end <= col or row_end <= row:
        return None  # the farm is outside this raster
    return Window(col, row, col_end - col, row_end - row)


def write_sparse_raster(src_path: Path, dst_path: Path, polygons) -> None:
    with rasterio.open(src_path) as src:
        profile = src.profile.copy()
        profile.update(
            driver="GTiff",
            tiled=True,
            blockxsize=512,
            blockysize=512,
            compress="deflate",
            sparse_ok=True,
        )
        windows = [w for w in (padded_window(src, p) for p in polygons) if w]
        with rasterio.open(dst_path, "w", **profile) as dst:
            for window in windows:
                dst.write(src.read(1, window=window), 1, window=window)


def build_fixtures(real_store: LayerStore, rows) -> None:
    index = real_store.read_index()
    assert index is not None
    polygons = farm_polygons(rows)

    shutil.rmtree(FIXTURE_MAPS_ROOT, ignore_errors=True)
    rasters_dir = FIXTURE_MAPS_ROOT / "layers" / "rasters"
    rasters_dir.mkdir(parents=True)
    for entry in index:
        src = real_store.raster_path(entry["raster_filename"])
        if is_lfs_pointer(src):
            raise SystemExit(f"{src} is a Git LFS pointer. Run 'git lfs pull' first.")
        write_sparse_raster(src, rasters_dir / entry["raster_filename"], polygons)
    (FIXTURE_MAPS_ROOT / "index.json").write_text(
        json.dumps(index, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )


def write_results_sheet(results, real_store: LayerStore) -> None:
    index = real_store.read_index()
    assert index is not None
    aliases = {}
    for entry in index:
        attributes = (
            real_store.read_attributes(entry["attributes_filename"], "es") or {}
        )
        aliases[str(entry["id"])] = attributes.get("alias") or f"Capa {entry['id']}"

    workbook = openpyxl.load_workbook(EXCEL_PATH)
    if RESULTS_SHEET in workbook.sheetnames:
        del workbook[RESULTS_SHEET]
    sheet = workbook.create_sheet(RESULTS_SHEET)
    sheet.append(
        [
            f"Deforestación por capa (umbral de superposición {OVERLAP_THRESHOLD_PERCENTAGE:g} %)."
        ]
    )
    layer_ids = sorted(aliases, key=int)
    sheet.append(
        ["ID", "Superficie [ha]", "Validación", *[aliases[i] for i in layer_ids]]
    )
    for farm_id, farm in results["farms"].items():
        values = []
        for layer_id in layer_ids:
            value = results["deforestation"][farm_id][layer_id]
            values.append("sin datos" if value is None else f"{value * 100:.2f} %")
        sheet.append(
            [
                farm_id,
                round(farm["area_m2"] / 10_000, 4),
                results["validation"]["status"][farm_id],
                *values,
            ]
        )
    sheet.append([])
    sheet.append(["Inconsistencias de validación"])
    for inconsistency in results["validation"]["inconsistencies"]:
        detail = inconsistency.get("reason") or (
            f"{inconsistency['percentage'] * 100:.2f} % ({inconsistency['criticality']})"
        )
        sheet.append(
            [inconsistency["type"], ", ".join(inconsistency["farmIds"]), detail]
        )
    sheet.column_dimensions["A"].width = 16
    for letter in "BCDEFGHI":
        sheet.column_dimensions[letter].width = 18
    workbook.save(EXCEL_PATH)


def main() -> None:
    real_store = get_layer_store()
    rows = read_farm_rows()

    real = run_pipeline(rows)
    build_fixtures(real_store, rows)

    set_layer_store(LayerStore(FIXTURE_MAPS_ROOT))
    try:
        fixture = run_pipeline(rows)
    finally:
        set_layer_store(None)
    if fixture != real:
        raise SystemExit("Sparse fixtures don't reproduce the real layers' results")

    EXPECTED_PATH.write_text(
        json.dumps(
            {
                "about": "Generated by `uv run python -m tests.regression.generate`.",
                "locale": LOCALE,
                "overlap_threshold_percentage": OVERLAP_THRESHOLD_PERCENTAGE,
                "results": real,
            },
            indent=2,
            ensure_ascii=False,
        )
        + "\n",
        encoding="utf-8",
    )
    write_results_sheet(real, real_store)
    size_kb = (
        sum(p.stat().st_size for p in FIXTURE_MAPS_ROOT.rglob("*") if p.is_file())
        / 1024
    )
    print(
        f"Wrote {EXPECTED_PATH.name}, fixtures ({size_kb:.0f} KB) and the Excel results sheet"
    )


if __name__ == "__main__":
    main()
